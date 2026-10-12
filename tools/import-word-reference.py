"""Build the app catalog from the supplied shared word reference.

Run with a Python installation that has openpyxl:
    python tools/import-word-reference.py
"""

import html
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "word-reference.xlsx"
OUTPUT = ROOT / "word-catalog.js"
LEVELS = {"A1", "A2", "B1", "B2", "C1", "C2"}


def split_values(value, separator):
    return list(dict.fromkeys(
        item.strip() for item in html.unescape(str(value or "")).split(separator) if item.strip()
    ))


sheet = load_workbook(SOURCE, read_only=True, data_only=True)["Справочник слов"]
rows = sheet.iter_rows(values_only=True)
header = next(rows)
if tuple(header[:5]) != ("Слово", "Перевод", "Тип", "Сложность", "Тема"):
    raise ValueError(f"Unexpected reference columns: {header!r}")

catalog = []
for number, row in enumerate(rows, start=2):
    padded = list(row) + [None] * max(0, 5 - len(row))
    word, translation, part_of_speech, cefr, topics = padded[:5]
    word = html.unescape(str(word or "")).strip()
    translations = split_values(translation, ",")
    cefr = str(cefr or "").strip().upper()
    if not word or not translations or not part_of_speech or cefr not in LEVELS:
        raise ValueError(f"Invalid word reference row {number}: {row!r}")
    catalog.append({
        "id": f"reference-{number - 1:05d}",
        "word": word,
        "translations": translations,
        "partOfSpeech": str(part_of_speech).strip(),
        "cefr": cefr,
        "topics": split_values(topics, ";"),
    })

OUTPUT.write_text(
    "// Generated from data/word-reference.xlsx by tools/import-word-reference.py.\n"
    "(() => {\n  'use strict';\n  window.VORDIK_WORD_CATALOG = "
    + json.dumps(catalog, ensure_ascii=False, separators=(",", ":"))
    + ";\n})();\n",
    encoding="utf-8",
)
print(f"Imported {len(catalog)} words into {OUTPUT.name}")
